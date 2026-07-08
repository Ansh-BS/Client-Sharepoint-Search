import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

os.environ.setdefault("SECRET_KEY", "test-secret")
os.environ.setdefault("STAFF_PASSWORD", "staffpw")
os.environ.setdefault("ADMIN_PASSWORD", "adminpw")

import openpyxl
import pytest


@pytest.fixture
def make_xlsx():
    """Build a spreadsheet matching the real file's layout.

    rows: list of (client_id, name, link_or_None). Column C gets the name
    as cell text and the link (if any) as the cell's hyperlink target.
    """
    def _make(path, rows, header=True):
        wb = openpyxl.Workbook()
        ws = wb.active
        if header:
            ws.append(["Client ID", "Particulars", "Sharepoint Path"])
        for cid, name, link in rows:
            ws.append([cid, name, name])
            if link:
                ws.cell(row=ws.max_row, column=3).hyperlink = link
        wb.save(path)
        return path
    return _make


@pytest.fixture
def data_path(tmp_path, monkeypatch):
    p = tmp_path / "clients.json"
    monkeypatch.setenv("CLIENTS_JSON", str(p))
    return p


@pytest.fixture
def client(data_path):
    from app import app
    app.config["TESTING"] = True
    with app.test_client() as c:
        yield c


@pytest.fixture
def logged_in(client):
    client.post("/login", data={"password": "staffpw"})
    return client
